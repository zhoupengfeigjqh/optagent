# -*- coding: utf-8 -*-
"""bom_core 纯逻辑单测。

运行方式（**宿主机本地**；容器只做部署与运行，不承担测试）：
    cd bom-service && python -m pytest -q tests

覆盖：URL 入参校验与受限下载、zip 内存解压（过滤/限额/同名冲突）、配置与 schema 自检、
S1.2/S1.5/S2/S2.4/S2.5 各条校验、递归展开与循环截断、累计用量连乘、按 schema 导出、
结果形状（标准 JSON）与产出命名（HTTP 均经 ``httpx.MockTransport`` 注入，不发真实请求）。

**不覆盖**：与真实BOM 的一致性比对——那是本服务职责之外，代码里也没有；
产出回写平台（``resultUrl``）在 ``tests/test_server_e2e.py``（真起服务 + 假回写端点）。
"""
import importlib
import io
import os
import zipfile

import httpx
import pytest
from openpyxl import Workbook, load_workbook

# ---------------------------------------------------------------- 夹具


def header_of(family="WCCMT"):
    first = "CODICE COMAU (SAP)" if family == "WCCMT" else "N°DISEGNO"
    return [first, "POSIZ.", "QTA'", "DESCRIZIONE", "MATERIALE   TABELLA    FORNITORE",
            "SIGLA    COMMERCIALE", "RICAMBIO", "REVIS.", "TAV."]


def xls_bytes(items, family="WCCMT", header=None):
    """构造一个最小子BOM：第 1 行标题行 + 明细行（列序与映射一致）。"""
    wb = Workbook()
    ws = wb.active
    ws.append(header if header is not None else header_of(family))
    for it in items:
        ws.append(it)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def item(code, posiz="1", qty=1, descr="件"):
    return [code, posiz, qty, descr, "STEEL", "D10", "", "A", "1"]


def zip_bytes(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


def mapping_of(families=("WCCMT", "WCCUN", "WCXXX", "VCPM")):
    """最小映射表：所有族共用同一套列名（够用且不依赖外部文件）。"""
    base = {"N°DISEGNO": "mat_code", "POSIZ.": "posiz", "QTA'": "qty", "DESCRIZIONE": "descr",
            "MATERIALE   TABELLA    FORNITORE": "materiale", "SIGLA    COMMERCIALE": "sigla",
            "RICAMBIO": "ricambio", "REVIS.": "revis", "TAV.": "tav"}
    base["CODICE COMAU (SAP)"] = "mat_code"
    return {"_header_detect": {"max_scan_rows": 12, "keywords": ["DISEGNO", "DRAWING", "CODICE"]},
            "_missing_component": {"patterns": [r"^WCCUN-\d+-S$", r"^WCXXX00000\d+$",
                                                r"^VCPM[A-Z0-9]+ZE$", r"^WCCMT-\d+$"]},
            **{f: dict(base) for f in families}}


def schema_of(extra: dict | None = None, order: list | None = None):
    """最小 schema：9 个映射字段 + 2 个计算生成 + 1 个人工维护；列序 = 17 列的子集。"""
    s = {
        "mat_code": {"col": 0, "name_cn": "物料编码", "source": "子BOM映射"},
        "posiz": {"col": 1, "name_cn": "位置号", "source": "子BOM映射"},
        "qty": {"col": 2, "name_cn": "单层用量", "source": "子BOM映射"},
        "total_qty": {"col": 3, "name_cn": "累计用量", "source": "计算生成"},
        "descr": {"col": 4, "name_cn": "名称", "source": "子BOM映射"},
        "materiale": {"col": 5, "name_cn": "材质", "source": "子BOM映射"},
        "sigla": {"col": 6, "name_cn": "规格", "source": "子BOM映射"},
        "ricambio": {"col": 7, "name_cn": "备件", "source": "子BOM映射"},
        "revis": {"col": 8, "name_cn": "版本", "source": "子BOM映射"},
        "tav": {"col": 9, "name_cn": "页号", "source": "子BOM映射"},
        "unit_attr": {"col": 10, "name_cn": "整件/散件属性", "source": "计算生成"},
        "unit_qty": {"col": 11, "name_cn": "整件/散件数量", "source": "人工维护"},
        "excel_column_order": ["mat_code", "posiz", "qty", "total_qty", "descr", "materiale",
                               "sigla", "ricambio", "revis", "tav", "unit_attr", "unit_qty"],
    }
    if extra:
        s.update(extra)
    if order is not None:
        s["excel_column_order"] = order
    return s


ROOT_ITEMS = [item("WCCUN-0001-S", "1", 2, "GR. COLONNE"), item("LEAF-0001", "2", 3, "叶子件")]
CHILD_ITEMS = [item("LEAF-0002", "1", 4, "下级叶子")]


def docs(root_items=None, child_items=None, root_name="WCCMT-00000B0570_00_0B.xls",
         extra_zip=None, schema=None, mapping=None):
    files = {"WCCUN-0001-S_00_0A.xls": xls_bytes(CHILD_ITEMS if child_items is None else child_items,
                                                 family="WCCUN")}
    files.update(extra_zip or {})
    return {
        "zip_bytes": zip_bytes(files),
        "total_bom_bytes": xls_bytes(ROOT_ITEMS if root_items is None else root_items),
        "total_bom_name": root_name,
        "schema_obj": schema if schema is not None else schema_of(),
        "mapping_obj": mapping if mapping is not None else mapping_of(),
    }


def codes(result):
    return [i["code"] for i in result["issues"]]


def load_core(monkeypatch, **env):
    """按给定环境变量重新加载 bom_core（该模块在导入期读取环境变量）。"""
    for k in ("BOM_MAX_INPUT_BYTES", "BOM_MAX_UNCOMPRESSED_BYTES", "BOM_MAX_FILES", "BOM_MAX_NODES"):
        monkeypatch.delenv(k, raising=False)
    for k, v in env.items():
        monkeypatch.setenv(k, str(v))
    import bom_core
    return importlib.reload(bom_core)


def transport(status: int, content=b""):
    return httpx.MockTransport(lambda request: httpx.Response(status, content=content))


# ---------------------------------------------------------------- 归一化 / 族


def test_norm_name_collapses_all_whitespace():
    core = load_core(pytest.MonkeyPatch())
    assert core.norm_name("SIGLA \n  COMMERCIALE") == "SIGLA COMMERCIALE"
    assert core.norm_name(None) == "None"


def test_family_of_uses_mapping_keys_verbatim():
    core = load_core(pytest.MonkeyPatch())
    fams = list(mapping_of())
    assert core.family_of("WCCMT-00000B0570_00_0B", fams) == "WCCMT"
    assert core.family_of("WCCUN-0000147248-S_00_0A", fams) == "WCCUN"
    assert core.family_of("VCPM6B2ER5ZE-B", fams) == "VCPM"
    assert core.family_of("WHAT-IS-THIS", fams) is None


def test_strip_version_suffix_only_for_root_inference():
    core = load_core(pytest.MonkeyPatch())
    assert core.strip_version_suffix("WCCMT-00000B0570_00_0B") == "WCCMT-00000B0570"
    assert core.strip_version_suffix("WCCUN-0000147248-S_00_0A") == "WCCUN-0000147248-S"
    # 索引 key 本身不截断：带 -A/-B 的 VCPM 原样保留
    assert core.strip_version_suffix("VCPM6B2ER5ZE-B") == "VCPM6B2ER5ZE-B"


# ---------------------------------------------------------------- URL / 下载


@pytest.mark.parametrize("url", [
    "", "not-a-url", "ftp://host/x.zip", "file:///etc/passwd", "临时空间/子BOM.zip", "数据准备/a.xls",
])
def test_non_http_inputs_are_rejected(monkeypatch, url):
    core = load_core(monkeypatch)
    assert core.check_http_url(url) is not None


@pytest.mark.parametrize("url", ["http://bom:8000/x", "https://mcp.example.com/a?p=共享空间/x.xlsx"])
def test_http_inputs_pass(monkeypatch, url):
    core = load_core(monkeypatch)
    assert core.check_http_url(url) is None


def test_download_reports_expired_signature(monkeypatch):
    core = load_core(monkeypatch)
    out = core.download("http://backend:3000/x", transport=transport(403))
    assert isinstance(out, str) and "签名可能已过期" in out


def test_download_rejects_oversized_body(monkeypatch):
    core = load_core(monkeypatch, BOM_MAX_INPUT_BYTES=10)
    out = core.download("http://backend:3000/x", transport=transport(200, b"x" * 100))
    assert isinstance(out, str) and "超过上限" in out


def test_download_returns_bytes(monkeypatch):
    core = load_core(monkeypatch)
    assert core.download("http://backend:3000/x", transport=transport(200, b"abc")) == b"abc"


def test_url_basename_prefers_signed_path_query(monkeypatch):
    core = load_core(monkeypatch)
    url = "http://backend:3000/api/files/raw?u=admin&p=%E6%95%B0%E6%8D%AE%E5%87%86%E5%A4%87/BOM%E6%8F%90%E4%BA%A4/%E5%AD%90BOM.zip"
    assert core.url_basename(url) == "子BOM.zip"
    assert core.url_basename("http://backend:3000/a/WCCMT-00000B0570_00_0B.xls") == "WCCMT-00000B0570_00_0B.xls"


# ---------------------------------------------------------------- S0 zip


def test_zip_skips_metadata_and_non_excel(monkeypatch):
    core = load_core(monkeypatch)
    issues = []
    data = zip_bytes({"__MACOSX/x.xls": b"junk", "~$tmp.xls": b"junk",
                      ".hidden.xls": b"junk", "readme.txt": b"junk",
                      "A_00_0A.xls": xls_bytes(ROOT_ITEMS)})
    members = core.read_zip_members(data, issues)
    assert [m[0] for m in members] == ["A_00_0A.xls"]
    assert issues == []


def test_zip_without_excel_reports_E003(monkeypatch):
    core = load_core(monkeypatch)
    issues = []
    core.read_zip_members(zip_bytes({"readme.txt": b"x"}), issues)
    assert [i["code"] for i in issues] == ["E003"]


def test_zip_duplicate_basename_reports_E010(monkeypatch):
    core = load_core(monkeypatch)
    issues = []
    data = zip_bytes({"a/b/A_00_0A.xls": xls_bytes(ROOT_ITEMS), "a/A_00_0A.xls": xls_bytes(ROOT_ITEMS)})
    members = core.read_zip_members(data, issues)
    assert [i["code"] for i in issues] == ["E010"]
    assert len(members) == 1


def test_broken_zip_reports_E002(monkeypatch):
    core = load_core(monkeypatch)
    issues = []
    core.read_zip_members(b"not a zip", issues)
    assert [i["code"] for i in issues] == ["E002"]


def test_zip_too_many_members_reports_E013(monkeypatch):
    core = load_core(monkeypatch, BOM_MAX_FILES=1)
    issues = []
    core.read_zip_members(zip_bytes({"A.xls": b"", "B.xls": b""}), issues)
    assert [i["code"] for i in issues] == ["E013"]


# ---------------------------------------------------------------- S0 schema 自检


def test_schema_order_missing_reports_E005a(monkeypatch):
    core = load_core(monkeypatch)
    issues = []
    core.check_schema_order({"mat_code": {"name_cn": "x"}}, issues)
    assert [i["code"] for i in issues] == ["E005a"]


def test_schema_order_undefined_and_duplicate(monkeypatch):
    core = load_core(monkeypatch)
    issues = []
    core.check_schema_order({"mat_code": {"col": 0, "name_cn": "物料编码"},
                             "excel_column_order": ["mat_code", "ghost", "mat_code"]}, issues)
    assert sorted(i["code"] for i in issues) == ["E005b", "E005c"]


def test_schema_defined_but_not_output_is_warn(monkeypatch):
    core = load_core(monkeypatch)
    issues = []
    s = schema_of()
    s["due_date"] = {"col": None, "name_cn": "交期(解析)", "source": "due_date_raw解析"}
    core.check_schema_order(s, issues)
    warn = [i for i in issues if i["code"] == "W005"]
    assert warn and "due_date" in warn[0]["message"]


def test_schema_col_mismatch_is_warn(monkeypatch):
    core = load_core(monkeypatch)
    issues = []
    core.check_schema_order({"a": {"col": 5, "name_cn": "A"}, "b": {"col": 1, "name_cn": "B"},
                             "excel_column_order": ["a", "b"]}, issues)
    assert [i["code"] for i in issues] == ["W007"]


def test_required_fields_are_derived_from_schema_source():
    """必需字段集由 schema 的 source 派生，不是硬编码 9 个。"""
    core = load_core(pytest.MonkeyPatch())
    issues = []
    cfg = core.load_config(schema_of({"brand": {"name_cn": "品牌", "source": "子BOM映射"}}),
                           mapping_of(), issues)
    assert "brand" in cfg["required_fields"]
    # 新字段没出现在任何族映射里 → 表头 1:1 检查必须报 E103
    d = docs(schema=schema_of({"brand": {"name_cn": "品牌", "source": "子BOM映射"}}))
    result = core.run_pipeline(**d)
    hit = [i for i in result["issues"] if i["code"] == "E103"]
    assert hit and "brand" in hit[0]["field"]


# ---------------------------------------------------------------- S1.2 / S1.5


def test_unknown_family_reports_E101():
    core = load_core(pytest.MonkeyPatch())
    d = docs(extra_zip={"WHATIS-1.xls": xls_bytes(ROOT_ITEMS)})
    result = core.run_pipeline(**d)
    assert "E101" in codes(result)


def test_header_not_found_reports_E102():
    core = load_core(pytest.MonkeyPatch())
    d = docs(root_items=[item("WCCUN-0001-S")])
    d["total_bom_bytes"] = xls_bytes([], header=["AA", "BB", "CC", "DD", "EE", "FF", "GG", "HH", "II"])
    result = core.run_pipeline(**d)
    assert "E102" in codes(result)
    issue = next(i for i in result["issues"] if i["code"] == "E102")
    assert "第 1~12 行" in issue["location"]


def test_missing_required_field_reports_E103_with_hint():
    core = load_core(pytest.MonkeyPatch())
    order_less = schema_of(order=["mat_code", "posiz", "qty", "descr", "materiale",
                                  "sigla", "ricambio", "revis", "total_qty", "unit_attr"])
    order_less["tav"] = {"col": 9, "name_cn": "页号", "source": "子BOM映射"}
    d = docs(schema=order_less)
    d["mapping_obj"]["WCCMT"].pop("TAV.")
    result = core.run_pipeline(**d)
    issue = next(i for i in result["issues"] if i["code"] == "E103")
    assert "tav" in issue["field"]
    assert "WCCMT" in issue["hint"]


def test_unmatched_column_reports_E104():
    core = load_core(pytest.MonkeyPatch())
    d = docs()
    hdr = header_of() + ["COLONNA STRANA"]
    d["total_bom_bytes"] = xls_bytes([item("WCCUN-0001-S")], header=hdr)
    result = core.run_pipeline(**d)
    issue = next(i for i in result["issues"] if i["code"] == "E104")
    assert "列9:COLONNA STRANA" in issue["evidence"]["unmatched"]


# ---------------------------------------------------------------- S2 / S2.4 / S2.5


def test_qty_not_numeric_is_warn_and_keeps_value():
    core = load_core(pytest.MonkeyPatch())
    # 注意：不用 "N/A" 这类字面量——pandas 默认会把它们读成 NaN（与旧实现同一行为）
    d = docs(child_items=[["LEAF-0002", "1", "abc", "下级叶子", "STEEL", "D10", "", "A", "1"]])
    result = core.run_pipeline(**d)
    assert "W301" in codes(result)
    assert "W501" in codes(result)               # 连乘随之失败
    assert result["status"] == "success"         # WARN 不改变状态
    node = next(n for n in result["_nodes"] if n["mat_code"] == "LEAF-0002")
    assert node["qty"] == "abc" and node["total_qty"] is None


def test_blank_mat_code_with_posiz_is_kept_as_leaf():
    """N°DISEGNO 空但 POSIZ. 有值 = 外购/标准件的合法明细行 → 必须保留成叶子行。

    这类行认不到子BOM 文件，也**不得**触发假 E202：空编码若进了编码认领队列，
    ``startswith("")`` 会匹配上所有文件（单文件时静默认错、多文件时报版本冲突）。
    """
    core = load_core(pytest.MonkeyPatch())
    d = docs(root_items=[item("", "6", 2, "BOCCOLA AUTOLUBRIFICANTE"), item("WCCUN-0001-S")])
    result = core.run_pipeline(**d, emit=True)
    assert result["status"] == "success"
    assert "E202" not in codes(result)
    blanks = [n for n in result["_nodes"] if n["mat_code"] == ""]
    assert len(blanks) == 1
    n = blanks[0]
    assert n["level"] == 1 and not n["is_component"] and not n["as_leaf"]
    assert (n["posiz"], n["qty"], n["descr"]) == ("6", 2, "BOCCOLA AUTOLUBRIFICANTE")
    assert n["total_qty"] == 2                       # 仍按父件连乘，量不丢
    assert result["_nodes"][1]["mat_code"] == ""     # 行序 = 源表行序（紧跟根行）
    # 产出里真的留了一行（mat_code 列空，posiz 原样带出）
    ws = load_workbook(io.BytesIO(result["_xlsx"]))["Total BOM"]
    blank_rows = [r for r in range(2, ws.max_row + 1) if not ws.cell(row=r, column=1).value]
    assert len(blank_rows) == 1
    assert ws.cell(row=blank_rows[0], column=2).value == "6"


def test_row_blank_in_both_code_and_posiz_is_skipped():
    """mat_code 与 posiz **都**空才是空行/填充行：照旧跳过，不产出节点。"""
    core = load_core(pytest.MonkeyPatch())
    d = docs(root_items=[item("WCCUN-0001-S", "1", 2, "GR. COLONNE"), [None] * 9,
                         item("LEAF-0001", "2", 3, "叶子件")])
    result = core.run_pipeline(**d)
    assert [n["mat_code"] for n in result["_nodes"]] == [
        "WCCMT-00000B0570_00_0B", "WCCUN-0001-S", "LEAF-0002", "LEAF-0001"]


def test_duplicate_row_reports_E201_with_positions():
    core = load_core(pytest.MonkeyPatch())
    d = docs(child_items=[item("LEAF-0002"), item("LEAF-0002")])
    result = core.run_pipeline(**d)
    issue = next(i for i in result["issues"] if i["code"] == "E201")
    assert "第1条 与 第2条" in issue["message"]


def test_code_matching_two_files_reports_E202():
    core = load_core(pytest.MonkeyPatch())
    d = docs(extra_zip={"WCCUN-0001-S_00_0B.xls": xls_bytes(CHILD_ITEMS, family="WCCUN")})
    result = core.run_pipeline(**d)
    issue = next(i for i in result["issues"] if i["code"] == "E202")
    assert issue["target"] == "WCCUN-0001-S"
    assert len(issue["evidence"]["candidates"]) == 2


def test_root_without_file_reports_E203():
    core = load_core(pytest.MonkeyPatch())
    result = core.run_pipeline(**docs(), root="WCCMT-NOT-EXIST")
    assert "E203" in codes(result)
    assert result["summary"]["stopped_at"] == "S2.5"


# ---------------------------------------------------------------- S4 / S5


def test_expand_shape_and_level1_unit_attr():
    core = load_core(pytest.MonkeyPatch())
    result = core.run_pipeline(**docs())
    nodes = result["_nodes"]
    # 根行 = 输入2 的文件 key（带版本后缀），与既有实现逐行一致
    assert [n["mat_code"] for n in nodes] == [
        "WCCMT-00000B0570_00_0B", "WCCUN-0001-S", "LEAF-0002", "LEAF-0001"]
    assert [n["level"] for n in nodes] == [0, 1, 2, 1]
    assert nodes[0]["is_component"] and not nodes[3]["is_component"]
    assert nodes[1]["unit_attr"] == "UNIT" and nodes[2]["unit_attr"] == ""
    assert nodes[0]["qty"] == 1 and nodes[0]["total_qty"] == 1


def test_total_qty_multiplies_along_the_tree():
    core = load_core(pytest.MonkeyPatch())
    result = core.run_pipeline(**docs(), order_units=2)
    by_code = {n["mat_code"]: n for n in result["_nodes"]}
    assert by_code["WCCUN-0001-S"]["total_qty"] == 4      # 2 × 2 台
    assert by_code["LEAF-0002"]["total_qty"] == 16        # 2 × 2 × 4


def test_calc_is_pure_three_layers():
    """2 × 3 × 4 = 24：不依赖 items_by_key，只用 nodes 自身。"""
    core = load_core(pytest.MonkeyPatch())
    nodes = [
        {"level": 0, "qty": 1, "mat_code": "R", "is_component": True, "source_key": "R"},
        {"level": 1, "qty": 2, "mat_code": "A", "is_component": True, "source_key": "R"},
        {"level": 2, "qty": 3, "mat_code": "B", "is_component": True, "source_key": "A"},
        {"level": 3, "qty": 4, "mat_code": "C", "is_component": False, "source_key": "B"},
    ]
    issues = []
    core.calc_total_qty(nodes, 1, issues)
    assert nodes[-1]["total_qty"] == 24 and issues == []


def test_cycle_is_truncated_with_warning():
    core = load_core(pytest.MonkeyPatch())
    # 子件反过来引用父件，构成环
    d = docs(child_items=[item("WCCMT-00000B0570", "1", 1, "环")])
    result = core.run_pipeline(**d)
    issue = next(i for i in result["issues"] if i["code"] == "W401")
    assert "循环引用截断" in issue["message"]
    assert result["status"] == "success"


def test_child_without_file_reports_W402_as_leaf():
    core = load_core(pytest.MonkeyPatch())
    d = docs(root_items=[item("WCCUN-9999999-S", "1", 1, "缺文件的组件")])
    result = core.run_pipeline(**d)
    issue = next(i for i in result["issues"] if i["code"] == "W402")
    assert "按叶子零件展开" in issue["message"]
    node = next(n for n in result["_nodes"] if n["mat_code"] == "WCCUN-9999999-S")
    assert node["as_leaf"] and not node["is_component"]


# ---------------------------------------------------------------- S6 导出


def test_export_follows_schema_column_order_and_blank_manual_columns():
    core = load_core(pytest.MonkeyPatch())
    order = ["mat_code", "descr", "qty", "total_qty", "unit_attr", "unit_qty"]
    result = core.run_pipeline(**docs(schema=schema_of(order=order)), emit=True)
    assert result["status"] == "success" and result["_xlsx"]
    wb = load_workbook(io.BytesIO(result["_xlsx"]))
    assert wb.sheetnames == ["Total BOM", "校验汇总"]
    ws = wb["Total BOM"]
    assert [c.value for c in ws[1]] == ["物料编码", "名称", "单层用量", "累计用量", "整件/散件属性", "整件/散件数量"]
    assert ws.max_column == len(order)
    # 人工维护列（unit_qty）恒空；组件行加粗；缩进 = level
    assert all(ws.cell(row=r, column=6).value is None for r in range(2, ws.max_row + 1))
    assert ws["A2"].font.bold is True and ws["A2"].alignment.indent == 0
    assert ws["A3"].alignment.indent == 1


def test_export_writes_issue_sheet_with_locators():
    core = load_core(pytest.MonkeyPatch())
    d = docs(root_items=[item("WCCUN-9999999-S", "1", 1, "缺文件的组件")])
    result = core.run_pipeline(**d, emit=True)
    ws = load_workbook(io.BytesIO(result["_xlsx"]))["校验汇总"]
    assert [c.value for c in ws[1]][:3] == ["错误码", "级别", "阶段"]
    assert [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)] == ["W402"]


def test_merge_failure_emits_no_xlsx():
    core = load_core(pytest.MonkeyPatch())
    d = docs(child_items=[item("LEAF-0002"), item("LEAF-0002")])
    result = core.run_pipeline(**d, emit=True)
    assert result["status"] == "failed" and result["_xlsx"] is None
    assert "E201" in codes(result)


# ---------------------------------------------------------------- 结果形状与产出命名


def test_result_status_driven_by_errors_only(monkeypatch):
    core = load_core(monkeypatch)
    warn = [core.issue("W301", "warn", "S2", "qty 非数值")]
    err = [core.issue("E103", "error", "S1.5", "缺字段")]
    assert core.build_result("m", {}, warn)["status"] == "success"
    assert core.build_result("m", {}, err)["status"] == "failed"
    assert core.build_result("m", {}, err + warn)["status"] == "failed"


def test_make_job_id_prefixed_and_unique(monkeypatch):
    core = load_core(monkeypatch)
    job = core.make_job_id(now=1_700_000_000.123)
    assert job.startswith("bom_1700000000123_") and len(job.rsplit("_", 1)[-1]) == 8
    assert len({core.make_job_id(now=1_700_000_000.0) for _ in range(50)}) == 50


def test_artifact_filename_carries_root_and_job_id(monkeypatch):
    """产出名 = 用户下载时看到的文件名：root 与 job_id 都在里面，一眼能对上是哪个任务。"""
    core = load_core(monkeypatch)
    assert core.artifact_filename("bom_1_ab", "WCCMT-00000B0570") == \
        "总BOM_WCCMT-00000B0570_bom_1_ab.xlsx"
    assert core.artifact_filename("bom_1_ab") == "总BOM_BOM_bom_1_ab.xlsx"     # root 缺省的兜底


def test_artifact_filename_sanitizes_root(monkeypatch):
    """root 来自文件名推断，可能带奇怪字符——它要进 URL 与磁盘名，必须洗成安全字符。"""
    core = load_core(monkeypatch)
    dirty = core.artifact_filename("bom_1_ab", "含/斜杠 与空格")
    assert "/" not in dirty and "\\" not in dirty and " " not in dirty
    assert dirty.startswith("总BOM_") and dirty.endswith("_bom_1_ab.xlsx")


def test_public_result_carries_artifact_only(monkeypatch):
    """对外结果：剥掉内部字段、补 artifact；**入口字段不在这一层拼**（回写成功后才由 server 补）。"""
    core = load_core(monkeypatch)
    name = "总BOM_X_bom_1_ab.xlsx"
    result = {"status": "success", "message": "合并完成", "_nodes": [1, 2], "_xlsx": b"PK"}
    out = core.public_result(result, filename=name, bytes_len=7, rows=322)
    assert "_nodes" not in out and "_xlsx" not in out                 # 内部字段不外泄
    assert out["artifact"] == {"filename": name, "rows": 322, "bytes": 7}
    # 路径只有平台回写响应才给得出；core 若自己编一个就是"点不开的入口"
    assert "filePath" not in out and "filePathName" not in out
    assert "viewUrl" not in out and "viewUrlName" not in out


def test_public_result_without_output_has_no_download_entry(monkeypatch):
    """没有产出（校验、或合并被阻断）时：artifact 为空，且**不留入口空字段**。"""
    core = load_core(monkeypatch)
    out = core.public_result({"status": "failed", "message": "校验未通过"})
    assert out["artifact"] is None
    assert "filePath" not in out and "filePathName" not in out


# ---------------------------------------------------------------- 全流程不落盘
# （`bom_core` 只算不写：回写是 `server.py` 在拿到产出的字节之后单独做的一步）


def test_pipeline_writes_no_files(monkeypatch, tmp_path):
    """服务侧零持久化：跑完一通（含导出）后工作目录不产生任何文件。"""
    core = load_core(monkeypatch)
    before = set(os.listdir(tmp_path))
    monkeypatch.chdir(tmp_path)
    result = core.run_pipeline(**docs(), emit=True)
    assert result["_xlsx"]
    assert set(os.listdir(tmp_path)) == before


def test_run_from_urls_reports_each_missing_input(monkeypatch):
    core = load_core(monkeypatch)

    def handler(request):
        return httpx.Response(404)

    result = core.run_from_urls(
        sub_bom_zip="http://b/x.zip", total_bom="http://b/y.xls",
        schema_url="http://b/s.json", mapping_url="http://b/m.json",
        transport=httpx.MockTransport(handler))
    assert result["status"] == "failed"
    assert [i["code"] for i in result["issues"]] == ["E001"] * 4
    # target 是**对外参数名**（相对路径口径），与工具签名一致，便于调用方对上入参
    assert {i["target"] for i in result["issues"]} == {
        "sub_bom_zip_path", "total_bom_path", "schema_path", "mapping_path"}
