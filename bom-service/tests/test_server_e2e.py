# -*- coding: utf-8 -*-
"""MCP 层端到端测试：真起服务、真走协议、**真把产出回写出去**。

链路：真实数据（BOM提交/）→ 本机静态服务铸成 http 直链 → 真实 ``server.py`` 进程
→ MCP 客户端 ``tools/call``（带平台注入的 ``resultUrl``）→ 服务算完 **POST 到假回写端点**
→ 结果里带回两条入口（``filePath`` = 报告、``mergedFile`` = xlsx），且端点收到的字节
一份是**能打开的 xlsx**、一份是**自包含的 HTML**。

（不需要真实数据的对应覆盖在 ``test_server_deliver.py`` 的 ``_deliver_outputs`` 用例里。）

这一层专门守住"注册得起来、参数齐、产出真能送出去"——``server.py`` 里
``@mcp.tool(structured_output=False)`` 与"工具内不得阻塞事件循环"那两条注释就是它要防的坑
（前者会让工具注册直接失败，后者会让并发调用互相卡死）。

只测 streamable-http：本服务只提供这一种传输。
"""
import asyncio
import functools
import http.server
import json
import os
import socket
import subprocess
import sys
import threading
import time
import urllib.parse
from pathlib import Path
from urllib.parse import quote

import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[2]
SERVICE = Path(__file__).resolve().parents[1]
SUBMIT = REPO / "BOM提交"
ZIP_NAME = "子BOM.zip"
TOTAL_BOM_NAME = "WCCMT-00000B0570_00_0B.xls"
SCHEMA_NAME = "总BOM标准字段.json"
MAPPING_NAME = "子BOM字段映射.json"

#: 真实数据展开行数（含根行），与 ``test_integration_real.EXPECTED_NODES`` 同源：
#: 含 5 行"mat_code 空、posiz 有值"的外购/标准件明细（曾经被整行丢掉）
EXPECTED_NODES = 327

pytestmark = pytest.mark.skipif(
    not (SUBMIT / ZIP_NAME).exists() or not (REPO / SCHEMA_NAME).exists(),
    reason="真实数据不在位，跳过端到端测试",
)


class _Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *args):        # 测试输出保持干净
        pass


def _serve(handler_cls):
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler_cls)
    srv.daemon_threads = True
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def _port_open(port: int) -> bool:
    with socket.socket() as s:
        s.settimeout(0.3)
        return s.connect_ex(("127.0.0.1", port)) == 0


def _free_port() -> int:
    # 缺省端口在本机可能已被别的程序占用（实测如此），不与之抢
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture(scope="module")
def static_url():
    srv = _serve(functools.partial(_Quiet, directory=str(REPO)))
    try:
        yield f"http://127.0.0.1:{srv.server_address[1]}"
    finally:
        srv.shutdown()


class _PutSink(http.server.BaseHTTPRequestHandler):
    """假的平台回写端点（``POST /api/files/put``）：收字节、记录请求、回 202 + ``path``。

    2026-10-22：一次合并会回写**两个**文件（xlsx + HTML 报告），所以除了 ``last`` 还要留
    全量 ``all``；``fail_html`` 用来单独打掉报告那一跳，验证"降级成只有数据文件"。
    """

    #: 最近一次请求（测试断言用）：``{"path":…, "query":{…}, "body":b"…", "headers":{…}}``
    last: dict = {}
    #: 全部请求（按到达顺序）
    all: list = []
    #: 收到的次数（失败场景断言"一次都没发"）
    calls = 0
    #: 返回状态码（换成 403 可测"回写被拒"）
    status = 202
    #: 只对 ``filename`` 以 ``.html`` 结尾的请求返回 403（测报告回写失败的降级）
    fail_html = False

    def log_message(self, *args):        # 测试输出保持干净
        pass

    def do_POST(self):
        length = int(self.headers.get("content-length") or 0)
        body = self.rfile.read(length)
        parsed = urllib.parse.urlsplit(self.path)
        query = {k: v[-1] for k, v in urllib.parse.parse_qs(parsed.query).items()}
        cls = type(self)
        cls.calls += 1
        entry = {"path": parsed.path, "query": query, "body": body,
                 "headers": dict(self.headers)}
        cls.last = entry
        cls.all.append(entry)
        if cls.status != 202 or (cls.fail_html
                                 and str(query.get("filename", "")).endswith(".html")):
            self.send_response(403 if cls.fail_html else cls.status)
            self.send_header("content-type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(b'{"error":{"code":"FILE_SIGN_INVALID"}}')
            return
        # 平台落盘名的形状：`临时空间/后台产出/{sid}_{服务给的文件名}`
        rel = f"临时空间/后台产出/{query.get('sid', 'u')}_{query.get('filename', 'x.xlsx')}"
        payload = json.dumps({"path": rel, "size": len(body)}, ensure_ascii=False).encode("utf-8")
        self.send_response(202)
        self.send_header("content-type", "application/json; charset=utf-8")
        self.send_header("content-length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)


@pytest.fixture()
def put_sink():
    """每个用例一个干净的假回写端点；返回 (base_url, 端点类)。"""
    handler = type("Sink", (_PutSink,), {"last": {}, "all": [], "calls": 0,
                                         "status": 202, "fail_html": False})
    srv = _serve(handler)
    try:
        yield f"http://127.0.0.1:{srv.server_address[1]}", handler
    finally:
        srv.shutdown()


@pytest.fixture(scope="module")
def mcp_server(tmp_path_factory):
    """真起 ``server.py``（streamable-http）。服务不落盘，故只需给一个空闲端口。"""
    port = _free_port()
    # 服务端日志必须**落文件**，绝不能挂一个没人读的管道：
    # uvicorn/httpx 每次请求都会打日志（Rich 还会把长 URL 折成十几行），管道写满后
    # `logger` 调用会阻塞事件循环 ⇒ 服务端整个卡死、响应永远发不出去。
    # 症状极具迷惑性：几乎不产生日志的用例（输入不可达）能过，一算就有大量日志的用例必卡。
    # 落文件既不阻塞，失败时还能把服务端日志打出来定位。
    log_path = tmp_path_factory.mktemp("mcp") / "server.log"
    log_file = open(log_path, "wb")
    env = {**os.environ, "BOM_MCP_PORT": str(port)}
    proc = subprocess.Popen([sys.executable, "-u", "server.py"], cwd=str(SERVICE), env=env,
                            stdout=log_file, stderr=subprocess.STDOUT)

    def server_log():
        return log_path.read_text(encoding="utf-8", errors="replace")

    try:
        deadline = time.time() + 30
        while time.time() < deadline and not _port_open(port):
            if proc.poll() is not None:
                pytest.fail(f"server.py 启动即退出：\n{server_log()[-2000:]}")
            time.sleep(0.3)
        if not _port_open(port):
            pytest.fail(f"server.py 30 秒内未监听 {port}\n{server_log()[-2000:]}")
        yield {"url": f"http://127.0.0.1:{port}/mcp",
               "base": f"http://127.0.0.1:{port}",
               "log": server_log}
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            proc.kill()
        log_file.close()


def _run_async(coro_factory, timeout: float):
    """在守护线程里跑一个协程并 join，保证"卡住 = 失败"，绝不让测试套件挂死。

    必须这么防：mcp 1.26 的 ``streamablehttp_client`` 在取消路径上会卡在会话终止
    （日志 ``Session termination failed``），实测一次清理要 5 分钟，远超预算。
    """
    box: dict = {}

    def worker():
        try:
            box["ok"] = asyncio.run(asyncio.wait_for(coro_factory(), timeout))
        except BaseException as e:                       # noqa: BLE001
            box["err"] = e

    t = threading.Thread(target=worker, daemon=True)
    t.start()
    t.join(timeout + 5)
    if t.is_alive():                                     # 清理都卡住：放弃，判失败
        raise TimeoutError(f"MCP 调用在 {timeout + 5:g} 秒内未返回（含清理）")
    if "err" in box:
        raise box["err"]
    return box["ok"]


def mcp_call(url: str, tool: str, arguments: dict, timeout: float = 30.0):
    """streamable-http 调用：建连 → 初始化 → 调工具，返回 (text, tool_names)。"""
    from mcp import ClientSession
    from mcp.client.streamable_http import streamablehttp_client

    async def run():
        async with streamablehttp_client(url) as (read, write, _):
            async with ClientSession(read, write) as session:
                await session.initialize()
                names = [t.name for t in (await session.list_tools()).tools]
                result = await session.call_tool(tool, arguments=arguments)
                return [c.text for c in result.content if getattr(c, "text", None)], names

    return _run_async(run, timeout)


def _url(static_url, rel_dir, name):
    path = f"{rel_dir}/{name}" if rel_dir else name
    return f"{static_url}/" + quote(path)


def file_args(static_url):
    """两个工具共有的四个文件参数（对外是相对路径口径；这里直接传铸造后的直链，模拟平台注入）。"""
    return {
        "sub_bom_zip_path": _url(static_url, "BOM提交", ZIP_NAME),
        "total_bom_path": _url(static_url, "BOM提交", TOTAL_BOM_NAME),
        "schema_path": _url(static_url, "", SCHEMA_NAME),
        "mapping_path": _url(static_url, "", MAPPING_NAME),
    }


def test_tools_registered_and_inputs_are_reported_when_unreachable(mcp_server):
    """工具注册成功；输入取不到时返回**标准 JSON 失败结果**（而不是 500/异常）。"""
    text, names = mcp_call(mcp_server["url"], "bom_validate", file_args("http://127.0.0.1:1"),
                           timeout=60)
    assert set(names) == {"bom_validate", "bom_merge"}
    assert text
    payload = json.loads(text[0])
    assert payload["status"] == "failed"
    assert {i["code"] for i in payload["issues"]} == {"E001"}
    assert {i["target"] for i in payload["issues"]} == {
        "sub_bom_zip_path", "total_bom_path", "schema_path", "mapping_path"}


def test_validate_returns_summary_and_no_output(mcp_server, static_url):
    """校验是纯计算工具：同一轮拿回完整汇总，**没有产出**（artifact 为空、也没有 filePath）。"""
    text, _ = mcp_call(mcp_server["url"], "bom_validate", file_args(static_url), timeout=60)
    payload = json.loads(text[0])
    assert payload["status"] == "success", payload["message"]
    assert payload["summary"]["tree"]["nodes"] == EXPECTED_NODES
    assert payload["artifact"] is None
    assert "filePath" not in payload and "filePathName" not in payload
    assert "viewUrl" not in payload and "viewUrlName" not in payload


def test_merge_writes_back_to_result_url(mcp_server, static_url, put_sink):
    """合并：**同一轮**返回统计 + 把 xlsx 与 HTML 报告各 POST 一次到「平台回写地址」，
    并把两次返回的 path 分别作为 ``mergedFile``（数据文件）与 ``filePath``（报告）。"""
    sink_url, sink = put_sink
    # 平台注入的地址都带原有一串 query（u/d/exp/sig/sid/call_id/tool）——服务必须一个不丢
    result_url = (f"{sink_url}/api/files/put?u=u1&d=%E4%B8%B4%E6%97%B6%E7%A9%BA%E9%97%B4"
                  f"&exp=1999999999999&sig=deadbeef&sid=th_1&call_id=c1&tool=bom__bom_merge")
    text, _ = mcp_call(mcp_server["url"], "bom_merge",
                       {**file_args(static_url), "resultUrl": result_url}, timeout=120)
    payload = json.loads(text[0])
    assert payload["status"] == "success", payload["message"]
    assert payload["summary"]["tree"]["nodes"] == EXPECTED_NODES
    # 结果统计就在这一次返回里（同步：不再有"受理时算不出来"这回事）
    assert payload["summary"]["counts"]["error"] == 0

    name = payload["artifact"]["filename"]
    assert name.startswith("总BOM_WCCMT-00000B0570_bom_") and name.endswith(".xlsx")
    assert payload["artifact"]["rows"] == EXPECTED_NODES and payload["artifact"]["bytes"] > 0

    # 两个文件、两条入口：顶层是报告（点开即读），mergedFile 是数据文件（点开即下载）
    assert sink.calls == 2
    report_name = payload["filePathName"]
    assert report_name == f"{name[:-len('.xlsx')]}_报告.html"
    assert payload["filePath"].startswith("临时空间/后台产出/")   # 平台返回的权威路径，原样搬运
    assert payload["filePath"].endswith(f"th_1_{report_name}")
    assert payload["mergedFile"]["filePath"].endswith(f"th_1_{name}")
    assert payload["mergedFile"]["filePathName"] == name
    assert "viewUrl" not in payload and "viewUrlName" not in payload

    # 第一次回写是 xlsx，第二次是报告；两次的原有 query 一个都不许丢
    xlsx_hit, report_hit = sink.all[0], sink.all[1]
    assert xlsx_hit["query"]["filename"] == name
    for hit in (xlsx_hit, report_hit):
        assert hit["query"]["u"] == "u1" and hit["query"]["sig"] == "deadbeef"
        assert hit["query"]["sid"] == "th_1" and hit["query"]["call_id"] == "c1"
        assert hit["query"]["tool"] == "bom__bom_merge"
        assert hit["query"]["summary"]           # 摘要必须有（界面标题的唯一来源）
        # 回写必须是"原样收字节"的类型：平台只给 octet-stream / text/plain 注册了 Buffer parser
        assert hit["headers"]["content-type"] == "application/octet-stream"
    wb = load_workbook(filename=__import__("io").BytesIO(xlsx_hit["body"]))
    assert wb.sheetnames == ["Total BOM", "校验汇总"]
    assert wb["Total BOM"].max_row == EXPECTED_NODES + 1      # 表头 + 数据行

    # 报告是自包含 HTML：无脚本（平台 sandbox="" 预览，脚本本来也不执行）
    assert report_hit["query"]["filename"] == report_name
    html = report_hit["body"].decode("utf-8")
    assert html.startswith("<!doctype html>") and "BOM 合并报告" in html
    assert "<details" in html and "<script" not in html.lower()
    assert name in html                        # 数据文件名写进正文（是文字，不是链接）


def test_merge_degrades_when_report_put_rejected(mcp_server, static_url, put_sink):
    """报告那一跳被拒：**降级**成只有数据文件——已经交付的入口绝不能被抹掉。"""
    sink_url, sink = put_sink
    sink.fail_html = True
    text, _ = mcp_call(mcp_server["url"], "bom_merge",
                       {**file_args(static_url),
                        "resultUrl": f"{sink_url}/api/files/put?u=u1&d=d&exp=1&sig=x&sid=th_1"},
                       timeout=120)
    payload = json.loads(text[0])
    assert payload["status"] == "success"
    assert sink.calls == 2                     # 报告那一跳发出过（被拒），不是没尝试
    name = payload["artifact"]["filename"]
    assert payload["filePath"].endswith(f"th_1_{name}")     # 顶层退回数据文件
    assert payload["filePathName"] == name
    assert "mergedFile" not in payload          # 同一路径不重复给两条入口
    assert "报告未交付" in payload["message"]


def test_merge_without_result_url_gives_no_file_path(mcp_server, static_url):
    """没有回写地址（直接调用服务）：结果照给，但**不给任何入口**——不编假路径。"""
    text, _ = mcp_call(mcp_server["url"], "bom_merge", file_args(static_url), timeout=120)
    payload = json.loads(text[0])
    assert payload["status"] == "success"
    assert payload["artifact"] is not None
    assert "filePath" not in payload and "filePathName" not in payload
    assert "未收到回写地址" in payload["message"]


def test_merge_blocked_by_error_has_no_output(mcp_server, static_url, put_sink):
    """有阻断性错误时不产出 BOM：status=failed、artifact 为空、**一次回写都不发**。"""
    sink_url, sink = put_sink
    args = {**file_args(static_url), "total_bom_sheet": "不存在的表",
            "resultUrl": f"{sink_url}/api/files/put?u=u1&d=d&exp=1&sig=x"}
    text, _ = mcp_call(mcp_server["url"], "bom_merge", args, timeout=120)
    payload = json.loads(text[0])
    assert payload["status"] == "failed"
    assert "E011" in {i["code"] for i in payload["issues"]}
    assert payload["artifact"] is None
    assert "filePath" not in payload
    assert sink.calls == 0


def test_merge_put_failure_does_not_fake_output(mcp_server, static_url, put_sink):
    """回写被拒（403）：**不假装有产出** —— 不给 filePath，原因写进 message。"""
    sink_url, sink = put_sink
    sink.status = 403
    text, _ = mcp_call(mcp_server["url"], "bom_merge",
                       {**file_args(static_url),
                        "resultUrl": f"{sink_url}/api/files/put?u=u1&d=d&exp=1&sig=x"},
                       timeout=120)
    payload = json.loads(text[0])
    assert sink.calls == 1
    assert "filePath" not in payload
    assert "产出未交付" in payload["message"] and "403" in payload["message"]
